#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""docx 表格解析的最小自检：**解析结果 vs 原文** 的列位对齐。

全链路此前没有任何「解析出来的东西还是不是原文」的判据 —— 业主实测到「开户银行/银行账号」
两行的值互换才发现。这里钉住的就是那一条：Word 的 layout grid（gridBefore/gridAfter，
「这行晚一格起步 / 早一格收尾」）不补空位的话，该行单元格整体左移，标签就跟邻行的值配上对。

纯 assert，不引测试框架。跑：  python tools/docx_grid_check.py
依赖 python-docx >= 1.2.0（grid_cols_before/after 是 1.2.0 才有的 API；
容器镜像的钉版在 docker/parser/Dockerfile，与 tools/requirements-embed.txt 须一致）。
"""
import os, sys, tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import embed_service

LABEL_BANK, VALUE_BANK = '开户银行', '中国光大银行股份有限公司长沙潇湘中路支行'
LABEL_ACCT, VALUE_ACCT = '银行账号', '79190188000132515'


def make_docx(path):
    """3 列网格；第二行 gridBefore=1（晚一格起步，只有 2 个 tc）—— 复现业主那份文档的形状。"""
    import docx
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 0).text, t.cell(0, 1).text = LABEL_BANK, VALUE_BANK   # 第三格留空
    t.cell(1, 0).text, t.cell(1, 1).text = LABEL_ACCT, VALUE_ACCT

    tr = t.rows[1]._tr
    tr.remove(tr.tc_lst[-1])                        # 去掉一个 tc，让这行只剩 2 格
    tr.get_or_add_trPr().get_or_add_gridBefore().val = 1   # 并声明「从第 2 列起步」
    doc.save(path)


def main():
    import docx
    assert hasattr(docx.table._Row, 'grid_cols_before'), \
        'python-docx 太老（需 >= 1.2.0）：没有 grid_cols_before，修不了错位'

    path = os.path.join(tempfile.mkdtemp(), 'grid.docx')
    make_docx(path)
    blocks, _, _ = embed_service._p_docx(path)

    tbl = [b['text'] for b in blocks if LABEL_ACCT in b['text']]
    assert len(tbl) == 1, f'表格块没解析出来：{blocks!r}'
    rows = [line.split(' | ') for line in tbl[0].split('\n')]
    assert len(rows) == 2, f'行数不对：{rows!r}'

    # 1) 每行都摊平到同样宽（3 = tblGrid 的列数）—— 空单元格保留占位，不许被 join 吞掉
    assert [len(r) for r in rows] == [3, 3], f'列位没对齐：{rows!r}'

    # 2) 标签与自己的值仍然相邻（错位时 银行账号 会落到第 0 列，跟上一行的值配对）
    for label, value in ((LABEL_BANK, VALUE_BANK), (LABEL_ACCT, VALUE_ACCT)):
        row = next(r for r in rows if label in r)
        assert row[row.index(label) + 1] == value, f'{label} 配错了值：{row!r}'

    # 3) 晚起步的那行确实从第 1 列开始（补空位补对了，不是碰巧两行都从 0 开始）
    assert rows[1][0] == '' and rows[1].index(LABEL_ACCT) == 1, f'gridBefore 没补：{rows[1]!r}'

    for r in rows:
        print(' | '.join(r))
    print('OK: docx 表格列位对齐')


if __name__ == '__main__':
    main()
